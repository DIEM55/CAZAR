"""Genera productos-investigados.xlsx a partir de datos.py (columnas de la skill)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import datos

COLS = ["Fecha", "Producto/Página", "Nicho", "Tipo de producto", "Score", "Confianza", "Veredicto",
        "Tier", "Antigüedad (días)", "Anuncios activos", "Países/idiomas", "Nº competidores clonando",
        "Saturación (sí/no)", "Precio detectado", "Plataforma de checkout", "Upsell (sí/no)",
        "Riesgo compliance", "Avatar", "Necesidad", "País con mejor potencial", "Tendencia",
        "Page ID", "Link de referencia"]

TIER_TXT = {"A": "A · cumple filtro duro", "B": "B · volumen ok, antigüedad probable",
            "C": "C · casi (6-19 días)", "D": "D · exótico, poco volumen"}

wb = Workbook()
ws = wb.active
ws.title = "Tracker"
ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="2F3B2F")
    c.alignment = Alignment(vertical="center", wrap_text=True)

for p in sorted(datos.P, key=lambda x: -x["score"]):
    ws.append([
        datos.FECHA, f'{p["producto"]} — {p["pagina"]}', p["nicho"], p["tipo"], p["score"],
        "Media" if p["tier"] in ("A", "B") else "Baja", p["veredicto"], TIER_TXT[p["tier"]],
        p["dias"], p["ads"], f'{p["idiomas"]} · {p["origen"]}', p["clones"],
        "sí" if p["saturado"] else "no", p["precio"], "No verificada (facebook.com y landings no accesibles)",
        "No verificado", p["compliance"], p["avatar"], p["necesidad"], p["pais"], "Primer chequeo",
        p["page_id"], p["links"][0],
    ])

widths = [11, 48, 22, 24, 7, 10, 16, 26, 10, 10, 30, 12, 10, 26, 24, 12, 30, 40, 40, 48, 14, 30, 50]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"

out = os.path.join(os.path.dirname(__file__), "productos-investigados.xlsx")
wb.save(out)
print("ok", out)
