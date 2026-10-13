"""Agrega al tracker los 20 productos en ascenso (cacería 13-oct-2026)."""
import os
from openpyxl import load_workbook
F = "2026-10-13"
L = "https://www.facebook.com/ads/library/?id={}"
NV = "No verificada (facebook.com no accesible)"
# producto, nicho, tipo, score, conf, ver, dias, ads, geo, clones, sat, pais, avatar, necesidad, page_id, ad_id, compliance
R = [
 ("Aprende a leer jugando 🦉 — Neurokids1 (red de 6 páginas)", "Infantil / lectura", "Pack lúdico de lectoescritura", 61, "Media", "GANADOR", 10.8, 38, "PL, IT, CZ, RO, GR, FR, ES (red ~92 ads)", 1, "no", "España, Chile, Uruguay, EE. UU. hispano", "Madres 28-45 con hijos de 4-7 años; abuelas 50-65 que regalan", "Aprender una habilidad (hijo que no arranca a leer)", "1125298484000687", "1630068381827850", "No"),
 ("160 recetas para tu Ninja Crispi — Cocina con Lola", "Cocina / electrodoméstico", "Recetario PDF", 56, "Media", "POTENCIAL", 6.9, 33, "ES, DE, EN, FR", 1, "no", "España, EE. UU. hispano, México, Chile", "Mujeres 35-60 que compraron o recibieron la Ninja Crispi", "Problema concreto: sacarle partido al aparato", "899211696601502", "1615453983318240", "Medio: marca registrada 'Ninja' en anuncios"),
 ("8 perros hiperrealistas al crochet — Crochet Pro", "Crochet / amigurumi", "Patrones PDF", 46, "Media", "POTENCIAL", 5.6, 26, "FR", 0, "no", "España, México, Chile, Argentina (en español no hay oferta)", "Mujeres 35-65 tejedoras con nivel medio, amantes de perros", "Aprender una habilidad / regalo personalizado", "1177673052104302", "1668371231538712", "No"),
 ("150 Boutique Gift Baskets — Berenice Garcia", "Negocio desde casa / eventos", "Libro de ideas (serie)", 55, "Media", "POTENCIAL", 7.7, 79, "EN (US)", 0, "no", "EE. UU., Canadá, Australia, Reino Unido; luego español", "Mujeres 35-60 que quieren un ingreso extra con regalos/eventos", "Ingresos extra (temporada de regalos)", "1087168021148527", "1443117621217189", "Bajo: evitar promesas de ingresos"),
 ("150 fichas de anatomía para colorear — Printora", "Estudio / salud", "Fichas imprimibles", 66, "Media", "GANADOR", 6.3, 34, "FR, IT, DE, EN", 2, "no", "Alemania, Países Bajos, Polonia, nórdicos (español ya saturado)", "Estudiantes de enfermería/medicina/veterinaria 18-35", "Aprender (memorizar anatomía)", "1232916199911605", "1124664126802508", "No"),
 ("500 actividades para adultos mayores — Digital Corner (+Bookbay SK)", "Cuidado de mayores", "Pack imprimible", 59, "Baja", "POTENCIAL", 3.0, 32, "HU, SK", 1, "no", "Alemania, Italia, España, Polonia, nórdicos", "Hijas 40-65 que cuidan a padres mayores; animadoras de residencias", "Problema concreto: entretener/estimular a un mayor", "1358074294058687", "1829682791389006", "No"),
 ("Ruhe im Kopf (sobrepensamiento) — Ruhe im Kopf", "Bienestar mental", "Guía (formato ?)", 61, "Baja", "GANADOR", 8.7, 38, "DE", 4, "no", "Alemania, Austria, Suiza", "Mujeres 35-55 con estrés y rumiación", "Problema concreto: calmar la mente", "1277187975474703", "1098550906254228", "Medio: claims de salud mental"),
 ("Explicado en dibujos (ebooks de estudio) — Étudefacile", "Estudio / formación", "Ebooks ilustrados", 50, "Media", "POTENCIAL", 6.8, 40, "FR (+IT/ES vía Studio Facile)", 1, "no", "Alemania, Países Bajos, Polonia", "Estudiantes 18-35 de oficios y salud", "Aprender más rápido", "1299578486577349", "1743486216949015", "No"),
 ("3 libros, 1 precio (niños: acoso y normas) — Digital Club Edu", "Crianza", "Pack de libros infantiles", 61, "Media", "GANADOR", 11.7, 54, "CS, RO, EL, PT(€)", 0, "no", "Polonia, Hungría, Alemania, España", "Madres 30-45 de niños 5-10", "Problema concreto: acoso escolar", "1361301167057013", "1419550166952042", "No"),
 ("1000 porqués (preguntas y respuestas) — Ezer Miért / Bin Neden", "Infantil / curiosidad", "Libro infantil", 49, "Media", "POTENCIAL", 4.9, 52, "HU, TR", 1, "no", "Polonia, Rumanía, España", "Padres 30-45 de niños 5-10", "Aprender", "1340954699103216", "1427285042193505", "No"),
 ("500.000+ diseños de sublimación — Tu Ebook Favorito", "Sublimación / POD", "Megapack de diseños", 51, "Baja", "POTENCIAL", 2.5, 160, "PL, RO, CS, PT", 0, "no", "Polonia, Chequia, Rumanía", "Emprendedoras 25-50 con sublimadora", "Ingresos extra", "1153819164474182", "1105926468862231", "Medio: derechos de diseños de terceros"),
 ("20 enigmas de detective para imprimir — Agenzia dei Piccoli Detective", "Infantil sin pantallas", "Imprimibles", 46, "Media", "POTENCIAL", 5.7, 15, "IT", 2, "no", "Italia, España (solo Mentes Mágicas)", "Madres 30-45 de niños 6-12", "Problema concreto: menos pantallas", "1440381472481078", "3700761470091822", "No"),
 ("150 experimentos de química sin laboratorio — Lauren Phillips", "Infantil / ciencia", "Guía de experimentos", 51, "Baja", "POTENCIAL", 1.5, 33, "EN", 2, "no", "EE. UU., Reino Unido, Australia (español ya competido)", "Madres homeschool 30-45", "Aprender", "1301771929686425", "977580018744582", "No"),
 ("200 patrones de corsé / 50 velas comestibles — Practical Guides", "Costura / velas", "Patrones + recetario", 63, "Media", "GANADOR", 8.4, 67, "EN", 3, "no", "EE. UU., Reino Unido, Alemania", "Mujeres 30-55 costureras/hobby", "Aprender una habilidad", "1280447521827796", "2271474530341702", "No"),
 ("300 problemas de ajedrez — The Knight's Notebook", "Ajedrez", "Cuadernos imprimibles", 54, "Baja", "POTENCIAL", 2.9, 22, "NO, IT, FR, EN", 0, "no", "Noruega, Dinamarca, Alemania, Países Bajos", "Padres y abuelos 35-70 de niños ajedrecistas", "Aprender", "1353979041129771", "1457546913012397", "No"),
 ("1.200 planos de acero para soldar / juguetes de madera — Mario Rossi", "Bricolaje / soldadura", "Pack de planos", 51, "Baja", "POTENCIAL", 3.5, 44, "NL, IT", 0, "no", "Países Bajos, Alemania, nórdicos", "Hombres 35-65 con taller en casa", "Ingresos extra / hobby", "1393906473803532", "1802993204071163", "No"),
 ("Cajas de chocolatería hechas en tu cocina — Gemas Digitales", "Repostería / emprender", "Curso/guía", 51, "Media", "POTENCIAL", 11.5, 41, "ES", 0, "no", "España, México, Chile", "Mujeres 30-55 que quieren vender bombones en Navidad", "Ingresos extra", "1129458703589779", "1130908282644296", "Bajo: evitar promesas de ingresos"),
 ("Guía de poses 1.200 fichas — Arte Digital", "Fotografía", "Fichas de poses", 66, "Media", "GANADOR", 5.9, 95, "FR, IT, EN, ES, JP", 1, "no", "Alemania, Países Bajos, Polonia", "Fotógrafas aficionadas 25-45", "Aprender", "276718215524267", "1117079470773438", "No"),
 ("Diario imprimible para fe complicada — Still Waters Journals", "Fe / journaling", "Diario imprimible", 46, "Baja", "POTENCIAL", 5.2, 22, "EN", 0, "no", "EE. UU., Canadá", "Mujeres cristianas 30-60 heridas por la iglesia", "Problema concreto: reconciliar la fe", "108018278680669", "1423484263182533", "No"),
 ("210+ patrones de marroquinería — Leathercraft Lab", "Cuero", "Patrones PDF", 51, "Baja", "POTENCIAL", 10.8, 20, "EN", 1, "no", "EE. UU., Alemania, nórdicos", "Hombres y mujeres 30-60 con hobby del cuero", "Aprender / ingresos extra", "927974923725307", "1841735747171829", "No"),
]
p = os.path.join(os.path.dirname(__file__), "productos-investigados.xlsx")
wb = load_workbook(p); ws = wb.active
have = {(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True)}
n = 0
for (prod, nicho, tipo, sc, conf, ver, d, ads, geo, cl, sat, pais, av, nec, pid, aid, comp) in R:
    if (F, prod) in have: continue
    ws.append([F, prod, nicho, tipo, sc, conf, ver, "Cacería en ascenso 13-oct", d, ads, geo, cl, sat, "No verificado", NV,
               "No verificado", comp, av, nec, pais, "Primer chequeo", pid, L.format(aid)]); n += 1
wb.save(p); print("filas:", n)
