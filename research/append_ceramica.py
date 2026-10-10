"""Agrega al tracker los competidores de 'plantillas cerámicas' (chequeo 10-oct-2026)."""
import os
from openpyxl import load_workbook

FECHA = "2026-10-10"
LINK = "https://www.facebook.com/ads/library/?id={}"
NICHO = "Cerámica / plantillas"
NO_CHECK = "No verificada (facebook.com y landings no accesibles)"
AVATAR = "Mujeres 30-60 que hacen cerámica de placa (slab) sin torno"
NECESIDAD = "Aprender una habilidad: cortar piezas a medida sin deformar ni desperdiciar arcilla"

# (producto/página, tipo, score, confianza, veredicto, días, ads, países/idiomas, clones, saturación,
#  precio, país potencial, tendencia, page_id, ad_id)
ROWS = [
    ("The Clay Blueprints", "Pack de plantillas imprimibles (EN)", 90, "Media", "GANADOR FUERTE", 25.8, 34,
     "EN · US/GB/DE/AU + Europa del Este", 10, "sí (en inglés)", "~US$42 (Dhs 154, según Excel unificado)",
     "EE. UU./Reino Unido; no recomendado para entrar ahora (mercado inglés con ~10 clones)",
     "→ estable (34 ads, igual que el Excel unificado del 7-oct)", "1343302345522178", "1692912678840741"),
    ("PrintClay Studio", "Plantillas pre-dimensionadas (EN/FR/NL/DA)", 90, "Media", "GANADOR FUERTE", 16.4, 33,
     "EN, FR, NL, DA · EE. UU. + Europa", 10, "sí (en inglés)", "No verificado",
     "Francia y Países Bajos: lo está traduciendo y testeando", "Primer chequeo", "1097177920156365", "1140551188635978"),
    ("La Biblia de las Plantillas Cerámicas — Ceramicaemprendedor", "Biblia de plantillas (ES/EN/DE)", 86, "Media",
     "GANADOR FUERTE", 16.8, 75, "ES, EN, DE · LATAM + España + EE. UU.", 5, "no",
     "No verificado", "España y México (es el líder hispano; página con 1005 anuncios activos, multi-producto)",
     "Primer chequeo", "860004577196412", "1414645386715189"),
    ("Kit de Plantillas Cerámicas — HI MERY", "Kit de plantillas + bonos", 71, "Media", "GANADOR", 8.8, 26,
     "ES · MX, ES, CL, CO, PE, UY, US (no pauta en AR)", 5, "no", "No verificado (anuncio '50% OFF + 7 bonos')",
     "Chile y España (mismo alcance que el producto del usuario, sin AR)", "Primer chequeo",
     "106784294503235", "1631515018383419"),
    ("ClayLab DIY / ClayCraft DIY (mismo operador)", "30+ plantillas imprimibles (EN/DE)", 68, "Media", "GANADOR",
     3.5, 21, "EN, DE · Europa", 10, "sí (en inglés)", "No verificado", "Alemania", "Primer chequeo",
     "978035102055668", "1397995229163204"),
    ("Pottery Blueprint", "Plantillas de papel para cerámica (EN)", 66, "Baja", "GANADOR", 16.0, 5,
     "EN · Europa (EUR)", 10, "sí (en inglés)", "No verificado", "Reino Unido/Irlanda", "Primer chequeo",
     "1339142019281477", "1134164862604217"),
    ("Clay Maker Lab", "Plantillas 'no wheel' (EN)", 58, "Baja", "POTENCIAL", 13.2, 8,
     "EN · Europa (EUR)", 10, "sí (en inglés)", "No verificado", "Europa del norte", "Primer chequeo",
     "1327089017144597", "1060421846815507"),
    ("Sistema Arcilla en Casa", "Método arcilla sin horno ni torno (ES)", 56, "Media", "POTENCIAL", 6.0, 5,
     "ES · LATAM + España + EE. UU.", 5, "no", "No verificado", "España", "Primer chequeo",
     "1353864484479328", "1808767807134186"),
    ("Barro Club", "+40 plantillas PDF para clases (ES)", 42, "Baja", "POTENCIAL", 3.1, 3,
     "ES · Argentina", 5, "no", "No verificado", "Argentina (profes de talleres)", "Primer chequeo",
     "1050393238168035", "2880014209043327"),
    ("Armá tu propia línea de cerámica — Ebooks del Futuro", "Ebook/kit emprendimiento cerámica (ES)", 41, "Baja",
     "POTENCIAL", 1.0, 9, "ES · Argentina", 5, "no", "No verificado", "Argentina", "Primer chequeo",
     "725759253957744", "938026992362764"),
    ("Plantillas para cerámica — Pulso cerámica", "Plantillas (ES, cobra en ARS)", 41, "Baja", "POTENCIAL", 3.0, 8,
     "ES · Argentina", 5, "no", "No verificado", "Argentina", "Primer chequeo",
     "269424112913360", "1760679795236946"),
]

path = os.path.join(os.path.dirname(__file__), "productos-investigados.xlsx")
wb = load_workbook(path)
ws = wb.active
existing = {(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True)}
added = 0
for (prod, tipo, sc, conf, ver, dias, ads, geo, clones, sat, precio, pais, tend, pid, aid) in ROWS:
    if (FECHA, prod) in existing:
        continue
    ws.append([FECHA, prod, NICHO, tipo, sc, conf, ver, "Chequeo cerámica 10-oct", dias, ads, geo, clones, sat,
               precio, NO_CHECK, "No verificado", "Bajo (sin claims de salud/ingresos detectados en títulos)",
               AVATAR, NECESIDAD, pais, tend, pid, LINK.format(aid)])
    added += 1
wb.save(path)
print("filas agregadas:", added)
