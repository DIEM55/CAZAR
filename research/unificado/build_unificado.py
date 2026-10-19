"""Une toda la investigación en un solo Excel: el unificado del usuario (07-oct),
el tracker propio (cerámica 10-oct + en ascenso 13-oct), el re-chequeo del 19-oct y el top 3."""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from recheck_19oct import RECHECK, TOP3

HERE = os.path.dirname(os.path.abspath(__file__))
SRC_USER = os.path.join(HERE, "fuente_usuario_2026-10-07.xlsx")
SRC_TRACKER = os.path.join(HERE, "..", "productos-investigados.xlsx")
OUT = os.path.join(HERE, "..", "productos-investigados-UNIFICADO-2026-10-19.xlsx")

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
TOP_FILL = PatternFill("solid", fgColor="E2EFDA")

COLS = ["Fuente", "Fecha", "Producto / Página", "Nicho", "Tipo de producto", "Score", "Confianza", "Veredicto",
        "Antigüedad (días)", "Anuncios activos", "Países / idiomas", "Nº competidores clonando", "Precio detectado",
        "Riesgo compliance", "País con mejor potencial", "Avatar", "Notas", "Page ID", "Link",
        "Cumple filtro duro (≥20 d y >15 anuncios)"]


def num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def filtro(dias, ads):
    d, a = num(dias), num(ads)
    if d is None or a is None:
        return "Sin dato"
    return "Sí" if d >= 20 and a > 15 else "No"


def style_sheet(ws, widths=None):
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


rows = []
# 1) Unificado del usuario (incluye el tracker del 06-oct, la lista maestra del 04-oct y los chequeos del 07-oct)
uws = load_workbook(SRC_USER, data_only=True)["Unificado"]
uh = [c.value for c in uws[1]]
for r in uws.iter_rows(min_row=2, values_only=True):
    if not any(r):
        continue
    d = dict(zip(uh, r))
    rows.append({
        "Fuente": d["Fuente"], "Fecha": d["Fecha"], "Producto / Página": d["Producto / Página"],
        "Nicho": d["Nicho"], "Tipo de producto": d["Tipo de producto"], "Score": num(d["Score"]),
        "Confianza": d["Confianza"], "Veredicto": d["Veredicto"], "Antigüedad (días)": num(d["Antigüedad (días)"]),
        "Anuncios activos": num(d["Anuncios activos"]), "Países / idiomas": d["Países / idiomas"],
        "Nº competidores clonando": d["Nº competidores clonando"], "Precio detectado": d["Precio detectado"],
        "Riesgo compliance": d["Riesgo compliance"], "País con mejor potencial": d["País con mejor potencial"],
        "Avatar": "", "Notas": d["Notas"], "Page ID": "", "Link": d["Link"],
    })

# Page IDs y avatar del tracker del 06-oct (mismo orden de productos que el unificado)
tws = load_workbook(SRC_TRACKER).active
th = [c.value for c in tws[1]]
trows = [dict(zip(th, r)) for r in tws.iter_rows(min_row=2, values_only=True)]
by_prod = {t["Producto/Página"]: t for t in trows if t["Fecha"] == "2026-10-06"}
for row in rows:
    t = by_prod.get(row["Producto / Página"])
    if t:
        row["Page ID"] = t["Page ID"] or ""
        row["Avatar"] = t["Avatar"] or ""

# 2) Tracker propio: cerámica (10-oct) y productos en ascenso (13-oct); el 06-oct ya está en el unificado
FUENTE = {"2026-10-10": "Chequeo cerámica 10/10", "2026-10-13": "Cacería en ascenso 13/10"}
for t in trows:
    if t["Fecha"] not in FUENTE:
        continue
    rows.append({
        "Fuente": FUENTE[t["Fecha"]], "Fecha": t["Fecha"], "Producto / Página": t["Producto/Página"],
        "Nicho": t["Nicho"], "Tipo de producto": t["Tipo de producto"], "Score": num(t["Score"]),
        "Confianza": t["Confianza"], "Veredicto": t["Veredicto"], "Antigüedad (días)": num(t["Antigüedad (días)"]),
        "Anuncios activos": num(t["Anuncios activos"]), "Países / idiomas": t["Países/idiomas"],
        "Nº competidores clonando": t["Nº competidores clonando"], "Precio detectado": t["Precio detectado"],
        "Riesgo compliance": t["Riesgo compliance"], "País con mejor potencial": t["País con mejor potencial"],
        "Avatar": t["Avatar"], "Notas": t["Tendencia"], "Page ID": t["Page ID"], "Link": t["Link de referencia"],
    })

for row in rows:
    row["Cumple filtro duro (≥20 d y >15 anuncios)"] = filtro(row["Antigüedad (días)"], row["Anuncios activos"])

wb = Workbook()

# Hoja 1: Top 3 caja rápida
ws = wb.active
ws.title = "Top 3 caja rápida 19-oct"
t3cols = ["#", "Producto a modelar", "Página de referencia", "Score 19-oct", "Anuncios activos hoy",
          "Tendencia vs. chequeo anterior", "Competencia en español", "Avatar", "Países para lanzar (orden)",
          "Precio sugerido", "Por qué hace caja rápido", "Cómo modelarlo", "Riesgos"]
ws.append(t3cols)
for t in TOP3:
    ws.append([t.get(c, "") for c in t3cols])
style_sheet(ws, {1: 4, 2: 30, 3: 26, 7: 30, 8: 30, 9: 28, 10: 22, 11: 45, 12: 45, 13: 30})

# Hoja 2: Re-chequeo 19-oct
ws = wb.create_sheet("Re-chequeo 19-oct")
rcols = ["Producto / Página", "Page ID", "Activos 19-oct", "Total histórico (ALL)", "Activos chequeo anterior",
         "Fecha chequeo anterior", "Tendencia", "Anuncio activo más viejo (días)", "Creados últimos 3 d / 4-7 d / antes",
         "Idioma / moneda / países", "Vendedores en español", "Score 19-oct", "Veredicto 19-oct", "Lectura"]
ws.append(rcols)
for r in sorted(RECHECK, key=lambda x: -(x.get("Score 19-oct") or 0)):
    ws.append([r.get(c, "") for c in rcols])
style_sheet(ws, {1: 34, 9: 22, 10: 26, 11: 32, 14: 50})

# Hoja 3: todo unificado, ordenado por score
ws = wb.create_sheet("Todo unificado")
ws.append(COLS + ["Ranking por score"])
rows.sort(key=lambda r: -(r["Score"] or 0))
for i, row in enumerate(rows, 1):
    ws.append([row[c] for c in COLS] + [i])
style_sheet(ws, {3: 40, 11: 26, 15: 30, 16: 30, 17: 40, 19: 30})

# Hoja 4: el top 3 anterior del usuario, tal cual
src = load_workbook(SRC_USER, data_only=True)["Top 3 y países"]
ws = wb.create_sheet("Top 3 anterior (07-oct)")
for r in src.iter_rows(values_only=True):
    if any(r):
        ws.append([v for v in r[:10]])

# Hoja 5: cerámica (producto propio)
ws = wb.create_sheet("Cerámica (tu producto)")
ws.append(["Tema", "Dato"])
for k, v in [
    ("Líder", "Ceramicaemprendedor (860004577196412), 'La Biblia de las Plantillas Cerámicas', en ~18 idiomas"),
    ("Idiomas donde sostiene anuncios viejos (rinde)", "ES, DE, FR, DK, NO, GR, HU, TR, CZ, PL"),
    ("Idiomas donde solo corre la última tanda (no rinde)", "IT, NL, SE, FI, RO, HR, JP, EN"),
    ("Países hispanos validados", "España, Chile, Uruguay (los líderes pautan ahí y sostienen)"),
    ("Apuestas baratas sin validar", "Puerto Rico, Panamá, Costa Rica: ningún líder los probó nunca"),
    ("Tus precios", "ES €19,90 · CL $14.990 · UY $549 · AR (lanzado 9-oct, 2 ventas con US$5 el día 1)"),
    ("Order bump", "Terrazo: sí. Lámparas de resina: no (marca 'Gema Academia')"),
    ("Locuciones", "research/locuciones_ceramica_ES_CL.txt / .docx (España 1-6, Chile 1-6)"),
]:
    ws.append([k, v])
style_sheet(ws, {1: 40, 2: 90})

wb.save(OUT)
print("guardado:", OUT, "| filas unificadas:", len(rows))
